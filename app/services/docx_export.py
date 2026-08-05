from pathlib import Path
import tempfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.schemas.tailoring import FormalResumeDocument
TAILORED_RESUME_EXPORT_DIR = Path(tempfile.gettempdir()) / "resume_tailor_exports"


INK = RGBColor(20, 27, 45)
ACCENT = RGBColor(43, 82, 168)
MUTED = RGBColor(94, 105, 128)
BODY_FONT = "Arial"
CJK_FONT = "Microsoft YaHei"


def export_formal_resume_docx(
    tailored_resume_id: str,
    formal_resume: FormalResumeDocument,
) -> Path:
    TAILORED_RESUME_EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    output_path = TAILORED_RESUME_EXPORT_DIR / f"{tailored_resume_id}.docx"

    document = Document()
    _configure_document(document)
    _add_resume_masthead(document, formal_resume)
    _add_resume_content(document, formal_resume)
    document.save(output_path)
    return output_path


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10)
    normal.font.color.rgb = INK
    normal._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.08

    heading = document.styles["Heading 1"]
    heading.font.name = BODY_FONT
    heading.font.size = Pt(11.5)
    heading.font.bold = True
    heading.font.color.rgb = ACCENT
    heading._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    heading._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    heading._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.space_after = Pt(4)
    heading.paragraph_format.keep_with_next = True

    bullet = document.styles["List Bullet"]
    bullet.font.name = BODY_FONT
    bullet.font.size = Pt(10)
    bullet._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    bullet._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    bullet._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    bullet.paragraph_format.left_indent = Inches(0.38)
    bullet.paragraph_format.first_line_indent = Inches(-0.2)
    bullet.paragraph_format.space_after = Pt(3)
    bullet.paragraph_format.line_spacing = 1.08

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    run = footer.add_run("ResumeTailor  |  ")
    _set_run_font(run, size=8, color=MUTED)
    _append_page_field(footer)


def _add_resume_masthead(
    document: Document,
    resume: FormalResumeDocument,
) -> None:
    name = (resume.name or "姓名").strip()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(2)
    title.paragraph_format.keep_with_next = True
    run = title.add_run(name)
    _set_run_font(run, size=22, color=INK, bold=True)

    if resume.headline:
        headline = document.add_paragraph()
        headline.paragraph_format.space_before = Pt(0)
        headline.paragraph_format.space_after = Pt(3)
        headline.paragraph_format.keep_with_next = True
        run = headline.add_run(resume.headline.strip())
        _set_run_font(run, size=11, color=ACCENT, bold=True)

    contact_items = [
        (
            f"{contact.label}：{contact.contact_value.strip()}"
            if contact.label and contact.label.strip()
            else contact.contact_value.strip()
        )
        for contact in resume.personal_contacts
        if contact.contact_value.strip()
    ]
    if not contact_items:
        contact_items = [
            item.strip()
            for item in (resume.email, resume.phone)
            if item and item.strip()
        ]
    if contact_items:
        contact = document.add_paragraph()
        contact.paragraph_format.space_before = Pt(0)
        contact.paragraph_format.space_after = Pt(8)
        run = contact.add_run("  |  ".join(contact_items))
        _set_run_font(run, size=9, color=MUTED)


def _add_resume_content(
    document: Document,
    resume: FormalResumeDocument,
) -> None:
    education_experiences = resume.education_experiences or resume.education
    _add_education_section(document, education_experiences)

    advantages = resume.advantages or resume.summary
    _add_text_section(document, "个人优势", advantages[:6], bullets=True)

    if resume.work_experiences:
        document.add_heading("工作经历", level=1)
        for work_experience in resume.work_experiences:
            heading = document.add_paragraph()
            heading.paragraph_format.space_before = Pt(3)
            heading.paragraph_format.space_after = Pt(2)
            heading.paragraph_format.keep_with_next = True
            company = heading.add_run(work_experience.company)
            _set_run_font(company, size=10.5, color=INK, bold=True)

            metadata = _work_metadata(work_experience)
            if metadata:
                meta_run = heading.add_run(f"  |  {metadata}")
                _set_run_font(meta_run, size=9, color=MUTED)

            for bullet in work_experience.bullets:
                _add_bullet(document, bullet)
    elif resume.experience:
        _add_text_section(
            document,
            "相关经历",
            resume.experience,
            bullets=True,
        )

    if resume.projects:
        document.add_heading("项目经历", level=1)
        for project in resume.projects:
            heading = document.add_paragraph()
            heading.paragraph_format.space_before = Pt(3)
            heading.paragraph_format.space_after = Pt(1)
            heading.paragraph_format.keep_with_next = True
            name_run = heading.add_run(project.name)
            _set_run_font(name_run, size=10.5, color=INK, bold=True)

            metadata = _project_metadata(project)
            if metadata:
                meta_run = heading.add_run(f"  |  {metadata}")
                _set_run_font(meta_run, size=9, color=MUTED)

            if project.technologies:
                technologies = document.add_paragraph()
                technologies.paragraph_format.space_before = Pt(0)
                technologies.paragraph_format.space_after = Pt(2)
                run = technologies.add_run(
                    "技术："
                    + " / ".join(
                        technology.strip()
                        for technology in project.technologies
                        if technology.strip()
                    )
                )
                _set_run_font(run, size=9, color=MUTED)

            for bullet in project.bullets:
                _add_bullet(document, bullet)

    if resume.honor_awards:
        document.add_heading("荣誉奖项", level=1)
        for honor_award in resume.honor_awards:
            heading = document.add_paragraph()
            heading.paragraph_format.space_before = Pt(3)
            heading.paragraph_format.space_after = Pt(2)
            heading.paragraph_format.keep_with_next = True
            name = heading.add_run(honor_award.name)
            _set_run_font(name, size=10.5, color=INK, bold=True)

            metadata = _honor_metadata(honor_award)
            if metadata:
                meta_run = heading.add_run(f"  |  {metadata}")
                _set_run_font(meta_run, size=9, color=MUTED)

            for bullet in honor_award.bullets:
                _add_bullet(document, bullet)

    related_skills = resume.related_skills or resume.skills
    if related_skills:
        document.add_heading("相关技能", level=1)
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run("  /  ".join(related_skills))
        _set_run_font(run, size=10, color=INK)


def _add_education_section(document: Document, education_experiences) -> None:
    if not education_experiences:
        return

    document.add_heading("教育经历", level=1)
    for education in education_experiences:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.keep_together = True
        school = paragraph.add_run(education.school)
        _set_run_font(school, size=10.5, color=INK, bold=True)

        details = [
            value.strip()
            for value in (education.degree, education.major)
            if value and value.strip()
        ]
        dates = _date_range(education.start_date, education.end_date)
        detail_text = " / ".join(details)
        suffix = "  |  ".join(
            value
            for value in (detail_text, dates)
            if value
        )
        if suffix:
            detail = paragraph.add_run(f"  |  {suffix}")
            _set_run_font(detail, size=9, color=MUTED)


def _add_text_section(
    document: Document,
    title: str,
    items: list[str],
    bullets: bool,
) -> None:
    normalized = [item.strip() for item in items if item.strip()]
    if not normalized:
        return

    document.add_heading(title, level=1)
    for item in normalized:
        if bullets:
            _add_bullet(document, item)
            continue

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.add_run(item)


def _add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.keep_together = True
    paragraph.add_run(text.strip())


def _project_metadata(project) -> str:
    values = []
    if project.role:
        values.append(project.role.strip())
    dates = _date_range(project.start_date, project.end_date)
    if dates:
        values.append(dates)
    return "  |  ".join(values)


def _work_metadata(work_experience) -> str:
    values = []
    if work_experience.job_title:
        values.append(work_experience.job_title.strip())
    dates = _date_range(
        work_experience.start_date,
        work_experience.end_date,
    )
    if dates:
        values.append(dates)
    return "  |  ".join(values)


def _honor_metadata(honor_award) -> str:
    return "  |  ".join(
        value.strip()
        for value in (honor_award.issuer, honor_award.date)
        if value and value.strip()
    )


def _date_range(start_date: str | None, end_date: str | None) -> str:
    if start_date and end_date:
        return f"{start_date.strip()} - {end_date.strip()}"
    return (start_date or end_date or "").strip()


def _set_run_font(
    run,
    size: float,
    color: RGBColor,
    bold: bool = False,
) -> None:
    run.font.name = BODY_FONT
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)


def _append_page_field(paragraph) -> None:
    run = paragraph.add_run()
    _set_run_font(run, size=8, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
