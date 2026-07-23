from pathlib import Path


SUPPORTED_EXTENSIONS = {".docx", ".pdf", ".txt", ".md"}


def read_document_text(file_path: Path) -> str:
    suffix = file_path.suffix.lower()

    if suffix not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported file type '{suffix}'. Supported: {supported}")

    if suffix in {".txt", ".md"}:
        return _read_plain_text(file_path)

    if suffix == ".docx":
        return _read_docx_text(file_path)

    if suffix == ".pdf":
        return _read_pdf_text(file_path)

    raise ValueError(f"Unsupported file type '{suffix}'")


def _read_plain_text(file_path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return file_path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue

    raise ValueError(f"Unable to decode text file: {file_path}")


def _read_docx_text(file_path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("Reading .docx files requires installing python-docx.") from exc

    document = Document(file_path)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
    table_cells = [
        cell.text.strip()
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]
    return "\n".join(text for text in paragraphs + table_cells if text)


def _read_pdf_text(file_path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("Reading .pdf files requires installing pypdf.") from exc

    reader = PdfReader(str(file_path))
    page_text = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(text.strip() for text in page_text if text.strip())
