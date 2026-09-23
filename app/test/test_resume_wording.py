import unittest

from docx import Document

from app.schemas.jds import ParsedJD
from app.schemas.resumes import ExperienceFact, ParsedResume, Project, Skill
from app.schemas.tailoring import (
    FactCheckReport, FormalResumeDocument, RequirementMatchReport,
    TailoredResumeDraft, TailoredSentence,
)
from app.services.docx_export import _add_resume_content
from app.services.resume_wording import resume_action_text
from app.services.tailoring import (
    _draft_needs_content_retry, assemble_formal_resume, build_rewrite_prompt,
    build_revision_prompt, rank_and_filter_draft, ranked_candidate_skill_sentences,
    validate_tailored_resume_draft,
)


class ResumeWordingTests(unittest.TestCase):
    def setUp(self):
        self.jd = ParsedJD(jd_id="jd", raw_text_length=0, required_skills=["Python"])
        self.resume = ParsedResume(
            resume_id="resume", name="张三",
            skills=[
                Skill(name=name, proficiency=level, evidence_fact_ids=[f"fact_{i}"])
                for i, (name, level) in enumerate([
                    ("Python", "熟练"), ("Java", "熟练"), ("C#", "熟悉"),
                    ("用户认证", "熟悉"), ("接口", "熟悉"), ("事件/委托", "熟悉"),
                ])
            ],
            experience_facts=[
                ExperienceFact(fact_id=f"fact_{i}", category="skill", entity_name=name,
                               fact_text=f"使用过{name}。")
                for i, name in enumerate(["Python", "Java", "C#", "用户认证", "接口", "事件/委托"])
            ],
            projects=[Project(
                project_id="project", name="应用后端", facts=[
                    ExperienceFact(fact_id="project_fact", category="project",
                                   entity_name="应用后端",
                                   fact_text="张三主要负责后端开发，参与用户认证模块实现。"),
                ],
            )],
        )
        self.match = RequirementMatchReport(jd_id="jd", resume_id="resume")

    def test_project_subject_removed_without_mutating_original_evidence(self):
        original = self.resume.model_dump()
        draft = TailoredResumeDraft(jd_id="jd", resume_id="resume")
        formal = assemble_formal_resume(self.jd, self.resume, draft)
        self.assertEqual(formal.name, "张三")
        self.assertEqual(formal.projects[0].bullets,
                         ["主要负责后端开发，参与用户认证模块实现。"])
        self.assertEqual(self.resume.model_dump(), original)
        self.assertEqual(resume_action_text("与张三协作开发。", "张三"), "与张三协作开发。")
        self.assertEqual(resume_action_text("张三的同事负责测试。", "张三"), "张三的同事负责测试。")
        self.assertEqual(resume_action_text("张三的项目主要以客户端开发为目标。", "张三"),
                         "项目主要以客户端开发为目标。")
        self.assertEqual(resume_action_text("Anna built it.", "Ann"), "Anna built it.")
        self.assertEqual(resume_action_text("张三负责开发。张三参与测试。", "张三"),
                         "负责开发。参与测试。")

    def test_fallback_groups_languages_preserves_levels_and_all_evidence(self):
        sentences = ranked_candidate_skill_sentences(self.jd, self.resume)
        languages = sentences[0]
        self.assertEqual(languages.sentence, "编程语言：熟练使用 Python、Java 进行开发；熟悉使用 C# 进行开发。")
        self.assertEqual(languages.source_fact_ids, ["fact_0", "fact_1", "fact_2"])
        concepts = sentences[1]
        self.assertEqual(concepts.sentence, "技术概念与实践：熟悉 用户认证、事件与委托。")
        self.assertEqual(concepts.source_fact_ids, ["fact_3", "fact_5"])
        self.assertNotIn("接口", "".join(item.sentence for item in sentences))

    def test_agent_grouped_sentences_survive_ranking_and_assembly(self):
        sentence = TailoredSentence(
            section="related_skills",
            sentence="编程语言：熟练使用 Python、Java 进行开发；熟悉 C#。",
            source_fact_ids=["fact_0", "fact_1", "fact_2"],
        )
        draft = TailoredResumeDraft(jd_id="jd", resume_id="resume", skills=[sentence])
        validated = validate_tailored_resume_draft(draft, self.jd, self.resume)
        ranked = rank_and_filter_draft(validated, self.match, self.jd, self.resume)
        self.assertEqual(ranked.skills, [sentence])
        self.assertEqual(assemble_formal_resume(self.jd, self.resume, ranked).related_skills,
                         [sentence.sentence])
        # Do not restore raw skills after a fact-check revision deletes them.
        draft.skills = []
        self.assertEqual(rank_and_filter_draft(draft, self.match, self.jd, self.resume).skills, [])

    def test_retry_allows_multiple_skills_per_sentence(self):
        payload = {
            "summary": [{"sentence": "优势"}] * 3,
            "skills": [{"sentence": "编程语言：熟练使用 Python、Java 进行开发；熟悉 C#。",
                        "source_fact_ids": ["fact_0", "fact_1", "fact_2"]},
                       {"sentence": "后端开发：熟悉用户认证；熟悉事件与委托。",
                        "source_fact_ids": ["fact_3", "fact_5"]}],
        }
        self.assertFalse(_draft_needs_content_retry(payload, self.resume))
        payload["skills"][0]["sentence"] = "熟练使用 Python / Java"
        self.assertTrue(_draft_needs_content_retry(payload, self.resume))

    def test_rewrite_and_revision_share_grouping_rules(self):
        draft = TailoredResumeDraft(jd_id="jd", resume_id="resume")
        prompts = [build_rewrite_prompt(self.jd, self.resume, self.match),
                   build_revision_prompt(self.jd, self.resume, draft,
                                         FactCheckReport(jd_id="jd", resume_id="resume"))]
        for prompt in prompts:
            self.assertIn("one coherent Chinese sentence per category", prompt)
            self.assertIn("用户认证", prompt)
            self.assertIn("never upgrade", prompt)
            self.assertIn("Never use slash", prompt)
            self.assertNotIn("one short phrase per skill", prompt)

    def test_export_keeps_skill_sentences_in_separate_bullets(self):
        skills = ["编程语言：熟悉 Python、Java。", "工程概念：熟悉面向对象编程。"]
        document = Document()
        _add_resume_content(document, FormalResumeDocument(related_skills=skills))
        paragraphs = [p for p in document.paragraphs if p.text in skills]
        self.assertEqual([p.text for p in paragraphs], skills)
        self.assertTrue(all(p.style.name == "List Bullet" for p in paragraphs))
        self.assertNotIn("/", "".join(p.text for p in document.paragraphs))


if __name__ == "__main__":
    unittest.main()
