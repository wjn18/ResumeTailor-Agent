import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.schemas.jds import JDRequirement, ParsedJD
from app.schemas.resumes import (
    Education,
    ExperienceFact,
    HonorAward,
    ParsedResume,
    Project,
    Skill,
    WorkExperience,
)
from app.schemas.tailoring import (
    FactCheckReport,
    FormalResumeUpdateRequest,
    RequirementMatchReport,
    TailoredResumeDraft,
    TailoredHonorAward,
    TailoredSentence,
    TailoredWorkExperience,
)
from app.services.document_reader import SUPPORTED_EXTENSIONS, read_document_text
from app.services.tailored_resume_storage import save_tailored_resume
from app.services.tailoring import (
    PARTIALLY_SUPPORTED,
    SUPPORTED,
    TailoringLLMClient,
    assemble_formal_resume,
    build_initial_tailored_resume,
    build_rewrite_prompt,
    build_tailored_resume,
    candidate_skills,
    collect_resume_facts,
    rank_and_filter_draft,
    review_tailored_resume,
    validate_fact_check_report,
    validate_match_report,
    validate_tailored_resume_draft,
)
from app.services.llm_client import (
    _education_experience_parse_incomplete,
    _honor_award_parse_incomplete,
    _personal_contact_parse_incomplete,
    _work_experience_parse_incomplete,
    build_resume_parse_prompt,
)


class AuditedTailoringClient(TailoringLLMClient):
    def __init__(self):
        self.fact_check_calls = 0
        self.revision_calls = 0

    def match_requirements(self, jd, resume):
        return {
            "jd_id": jd.jd_id,
            "resume_id": resume.resume_id,
            "matches": [
                {
                    "requirement_id": "req_001",
                    "requirement_text": "FastAPI",
                    "match_status": "matched",
                    "matched_fact_ids": ["fact_001"],
                    "reasoning": "The cited fact explicitly names FastAPI.",
                }
            ],
        }

    def rewrite_resume(self, jd, resume, match_report):
        return {
            "jd_id": jd.jd_id,
            "resume_id": resume.resume_id,
            "headline": jd.job_title,
            "summary": [
                {
                    "section": "advantages",
                    "sentence": "使用 FastAPI 服务百万用户。",
                    "source_fact_ids": ["fact_001"],
                }
            ],
            "work_experiences": [
                {
                    "work_experience_id": "work_001",
                    "bullets": [
                        {
                            "section": "work_experience",
                            "sentence": "使用 FastAPI 构建后端服务。",
                            "source_fact_ids": ["fact_001"],
                        }
                    ],
                }
            ],
            "honor_awards": [
                {
                    "honor_award_id": "honor_001",
                    "bullets": [
                        {
                            "section": "honor_award",
                            "sentence": "获得校级优秀毕业设计。",
                            "source_fact_ids": ["fact_honor_001"],
                        }
                    ],
                }
            ],
            "skills": [],
        }

    def fact_check_resume(self, jd_id, resume, draft):
        self.fact_check_calls += 1
        sentence = draft.summary[0]
        is_revised = sentence.sentence == "使用 FastAPI 构建后端服务。"
        checks = [
            {
                "section": sentence.section,
                "sentence": sentence.sentence,
                "source_fact_ids": sentence.source_fact_ids,
                "support_status": (
                    SUPPORTED if is_revised else PARTIALLY_SUPPORTED
                ),
                "issue": None if is_revised else "百万用户没有事实依据。",
                "suggestion": None if is_revised else "删除未提供的用户规模。",
            }
        ]
        for work_experience in draft.work_experiences:
            for bullet in work_experience.bullets:
                checks.append(
                    {
                        "section": bullet.section,
                        "sentence": bullet.sentence,
                        "source_fact_ids": bullet.source_fact_ids,
                        "support_status": SUPPORTED,
                        "issue": None,
                        "suggestion": None,
                    }
                )
        for honor_award in draft.honor_awards:
            for bullet in honor_award.bullets:
                checks.append(
                    {
                        "section": bullet.section,
                        "sentence": bullet.sentence,
                        "source_fact_ids": bullet.source_fact_ids,
                        "support_status": SUPPORTED,
                        "issue": None,
                        "suggestion": None,
                    }
                )
        for skill in draft.skills:
            checks.append(
                {
                    "section": skill.section,
                    "sentence": skill.sentence,
                    "source_fact_ids": skill.source_fact_ids,
                    "support_status": SUPPORTED,
                    "issue": None,
                    "suggestion": None,
                }
            )
        return {
            "jd_id": jd_id,
            "resume_id": resume.resume_id,
            "checks": checks,
        }

    def revise_after_fact_check(self, jd, resume, draft, fact_check_report):
        self.revision_calls += 1
        return {
            "jd_id": jd.jd_id,
            "resume_id": resume.resume_id,
            "headline": jd.job_title,
            "summary": [
                {
                    "section": "advantages",
                    "sentence": "使用 FastAPI 构建后端服务。",
                    "source_fact_ids": ["fact_001"],
                }
            ],
            "work_experiences": [
                item.model_dump()
                for item in draft.work_experiences
            ],
            "honor_awards": [
                item.model_dump()
                for item in draft.honor_awards
            ],
            "skills": [],
        }


class TwoPassRevisionClient(AuditedTailoringClient):
    def fact_check_resume(self, jd_id, resume, draft):
        report = super().fact_check_resume(jd_id, resume, draft)
        if self.revision_calls < 2:
            report["checks"][0].update(
                {
                    "support_status": PARTIALLY_SUPPORTED,
                    "issue": "The first correction still needs another audit pass.",
                    "suggestion": "Rewrite once more using only the cited fact.",
                }
            )
        return report


def sample_jd() -> ParsedJD:
    return ParsedJD(
        jd_id="jd_test",
        company="示例科技",
        job_title="后端工程师",
        raw_text_length=30,
        requirements=[
            JDRequirement(
                requirement_id="req_001",
                category="required_skill",
                requirement_text="FastAPI",
                keywords=["FastAPI"],
            )
        ],
    )


def sample_resume() -> ParsedResume:
    return ParsedResume(
        resume_id="resume_test",
        name="王建",
        email="wang@example.com",
        phone="13800000000",
        education=[
            Education(
                school="示例大学",
                degree="本科",
                major="计算机科学",
                start_date="2019",
                end_date="2023",
            )
        ],
        skills=[
            Skill(name="Python", proficiency="熟悉"),
            Skill(
                name="FastAPI",
                proficiency="熟练",
                evidence_fact_ids=["fact_001"],
            ),
        ],
        work_experiences=[
            WorkExperience(
                work_experience_id="work_001",
                company="示例科技",
                job_title="后端工程师",
                start_date="2023.06",
                end_date="2025.06",
                facts=[
                    ExperienceFact(
                        fact_id="fact_001",
                        category="work",
                        entity_name="示例科技",
                        fact_text="使用 FastAPI 构建后端服务。",
                    )
                ],
            )
        ],
        honor_awards=[
            HonorAward(
                honor_award_id="honor_001",
                name="优秀毕业设计",
                issuer="示例大学",
                date="2023.06",
                facts=[
                    ExperienceFact(
                        fact_id="fact_honor_001",
                        category="honor_award",
                        entity_name="优秀毕业设计",
                        fact_text="获得校级优秀毕业设计。",
                    )
                ],
            )
        ],
        projects=[
            Project(
                project_id="project_001",
                name="简历生成系统",
                role="后端开发",
                technologies=["Python", "FastAPI"],
                facts=[
                    ExperienceFact(
                        fact_id="fact_project_001",
                        category="project",
                        entity_name="简历生成系统",
                        fact_text="实现简历解析与存储流程。",
                    )
                ],
            )
        ],
    )


class TailoringWorkflowTests(unittest.TestCase):
    def test_prompts_require_atomic_facts_summaries_and_skill_sentences(self):
        parse_prompt = build_resume_parse_prompt("使用 FastAPI 开发接口。")
        rewrite_prompt = build_rewrite_prompt(
            sample_jd(),
            sample_resume(),
            RequirementMatchReport(
                jd_id="jd_test",
                resume_id="resume_test",
            ),
        )

        self.assertIn("atomic facts", parse_prompt)
        self.assertIn('"proficiency": "了解 | 熟悉 | 熟练 | 精通"', parse_prompt)
        self.assertIn("Return 3-6 concise Chinese bullet sentences", rewrite_prompt)
        self.assertIn("work_experience_id", parse_prompt)
        self.assertIn("honor_award_id", parse_prompt)
        self.assertIn("education_experience_id", parse_prompt)
        self.assertIn("personal_contact_id", parse_prompt)
        self.assertIn("Never invent or rewrite company", rewrite_prompt)
        self.assertIn("Never invent or rewrite the award name", rewrite_prompt)
        self.assertIn("Use the stored proficiency exactly", rewrite_prompt)
        self.assertIn(
            "A skill must still appear here when it is already mentioned",
            rewrite_prompt,
        )
        self.assertIn(
            "finally broadly useful skills such as Excel",
            rewrite_prompt,
        )
        self.assertIn('"proficiency": "熟练"', rewrite_prompt)

    def test_candidate_skills_keep_proficiency_and_evidence(self):
        skills = candidate_skills(sample_resume())
        fastapi = next(skill for skill in skills if skill["name"] == "FastAPI")

        self.assertEqual(fastapi["proficiency"], "熟练")
        self.assertEqual(fastapi["evidence_fact_ids"], ["fact_001"])

    def test_work_section_triggers_structured_parse_retry_check(self):
        self.assertTrue(
            _work_experience_parse_incomplete(
                {"work_experiences": []},
                "工作经历\n示例科技 后端工程师",
            )
        )
        self.assertFalse(
            _work_experience_parse_incomplete(
                {
                    "work_experiences": [
                        {
                            "company": "示例科技",
                            "work_experience_id": "work_001",
                        }
                    ]
                },
                "工作经历\n示例科技 后端工程师",
            )
        )

    def test_honor_section_triggers_structured_parse_retry_check(self):
        self.assertTrue(
            _honor_award_parse_incomplete(
                {"honor_awards": []},
                "荣誉奖项\n校级优秀毕业设计",
            )
        )
        self.assertFalse(
            _honor_award_parse_incomplete(
                {
                    "honor_awards": [
                        {
                            "name": "优秀毕业设计",
                            "honor_award_id": "honor_001",
                            "facts": [{"fact_id": "fact_honor_001"}],
                        }
                    ]
                },
                "荣誉奖项\n校级优秀毕业设计",
            )
        )

    def test_profile_sections_trigger_structured_parse_retry_checks(self):
        self.assertTrue(
            _education_experience_parse_incomplete(
                {"education_experiences": []},
                "教育经历\n示例大学 计算机科学本科",
            )
        )
        self.assertTrue(
            _personal_contact_parse_incomplete(
                {"personal_contacts": []},
                "邮箱：wang@example.com",
            )
        )

    def test_profile_metadata_is_promoted_to_fact_types(self):
        resume = sample_resume()

        self.assertEqual(
            resume.education_experiences[0].facts[0].category,
            "education_experience",
        )
        self.assertEqual(
            {contact.contact_type for contact in resume.personal_contacts},
            {"email", "phone"},
        )
        collected_fact_ids = {
            fact.fact_id
            for fact in collect_resume_facts(resume)
        }
        self.assertIn(
            resume.education_experiences[0].facts[0].fact_id,
            collected_fact_ids,
        )
        self.assertTrue(
            all(
                contact.facts[0].fact_id not in collected_fact_ids
                for contact in resume.personal_contacts
            )
        )

    def test_work_bullet_cannot_cite_a_project_fact(self):
        draft = TailoredResumeDraft(
            jd_id="jd_test",
            resume_id="resume_test",
            work_experiences=[
                TailoredWorkExperience(
                    work_experience_id="work_001",
                    bullets=[
                        TailoredSentence(
                            section="work_experience",
                            sentence="实现简历解析与存储流程。",
                            source_fact_ids=["fact_project_001"],
                        )
                    ],
                )
            ],
        )

        with self.assertRaisesRegex(ValueError, "valid source_fact_id"):
            validate_tailored_resume_draft(
                draft,
                sample_jd(),
                sample_resume(),
            )

    def test_honor_bullet_cannot_cite_a_project_fact(self):
        draft = TailoredResumeDraft(
            jd_id="jd_test",
            resume_id="resume_test",
            honor_awards=[
                TailoredHonorAward(
                    honor_award_id="honor_001",
                    bullets=[
                        TailoredSentence(
                            section="honor_award",
                            sentence="实现简历解析与存储流程。",
                            source_fact_ids=["fact_project_001"],
                        )
                    ],
                )
            ],
        )

        with self.assertRaisesRegex(ValueError, "valid source_fact_id"):
            validate_tailored_resume_draft(
                draft,
                sample_jd(),
                sample_resume(),
            )

    def test_formal_resume_uses_only_generated_skill_sentences(self):
        jd = sample_jd()
        resume = sample_resume()
        draft = TailoredResumeDraft(
            jd_id=jd.jd_id,
            resume_id=resume.resume_id,
            skills=[
                TailoredSentence(
                    section="skills",
                    sentence="熟练使用 FastAPI，能够完成后端接口开发。",
                    source_fact_ids=["fact_001"],
                )
            ],
        )

        formal_resume = assemble_formal_resume(jd, resume, draft)

        self.assertEqual(
            formal_resume.related_skills,
            ["熟练使用 FastAPI，能够完成后端接口开发。"],
        )
        self.assertNotIn("Python", formal_resume.related_skills)
        self.assertNotIn("FastAPI", formal_resume.related_skills)

    def test_advantages_are_ranked_without_removing_related_skills(self):
        advantages = [
            TailoredSentence(
                section="advantages",
                sentence=f"优势 {index}",
                source_fact_ids=[f"fact_{index}"],
            )
            for index in range(1, 8)
        ]
        draft = TailoredResumeDraft(
            jd_id="jd_test",
            resume_id="resume_test",
            summary=advantages,
            skills=[
                TailoredSentence(
                    section="related_skills",
                    sentence="重复技能",
                    source_fact_ids=["fact_7"],
                ),
                TailoredSentence(
                    section="related_skills",
                    sentence="补充技能",
                    source_fact_ids=["fact_other"],
                ),
            ],
        )
        report = RequirementMatchReport(
            jd_id="jd_test",
            resume_id="resume_test",
            matches=[
                {
                    "requirement_id": "req_001",
                    "requirement_text": "核心要求",
                    "match_status": "matched",
                    "matched_fact_ids": ["fact_7"],
                    "reasoning": "最高适配",
                }
            ],
        )

        result = rank_and_filter_draft(draft, report)

        self.assertEqual(len(result.summary), 6)
        self.assertEqual(result.summary[0].sentence, "优势 7")
        self.assertEqual(
            [item.sentence for item in result.skills],
            ["重复技能", "补充技能"],
        )

    def test_related_skills_include_advantage_skill_and_general_skill_last(self):
        resume = sample_resume().model_copy(
            update={
                "skills": [
                    Skill(
                        name="Excel",
                        proficiency="熟悉",
                        evidence_fact_ids=["fact_excel"],
                    ),
                    Skill(
                        name="FastAPI",
                        proficiency="熟练",
                        evidence_fact_ids=["fact_001"],
                    ),
                ],
                "experience_facts": [
                    ExperienceFact(
                        fact_id="fact_excel",
                        category="skill",
                        entity_name="Excel",
                        fact_text="日常使用 Excel 整理项目数据。",
                    )
                ],
            }
        )
        draft = TailoredResumeDraft(
            jd_id="jd_test",
            resume_id=resume.resume_id,
            summary=[
                TailoredSentence(
                    section="advantages",
                    sentence="熟练使用 FastAPI 开发后端接口。",
                    source_fact_ids=["fact_001"],
                )
            ],
        )
        report = RequirementMatchReport(
            jd_id="jd_test",
            resume_id=resume.resume_id,
            matches=[
                {
                    "requirement_id": "req_001",
                    "requirement_text": "FastAPI",
                    "match_status": "matched",
                    "matched_fact_ids": ["fact_001"],
                    "reasoning": "岗位明确要求 FastAPI。",
                }
            ],
        )

        result = rank_and_filter_draft(
            draft,
            report,
            sample_jd(),
            resume,
        )

        self.assertEqual(
            [item.sentence for item in result.skills],
            ["熟练使用 FastAPI", "熟悉使用 Excel"],
        )
        self.assertEqual(
            result.skills[0].source_fact_ids,
            ["fact_001"],
        )

    def test_formal_work_experience_preserves_parsed_metadata(self):
        resume = sample_resume()
        draft = TailoredResumeDraft(
            jd_id="jd_test",
            resume_id=resume.resume_id,
            work_experiences=[
                TailoredWorkExperience(
                    work_experience_id="work_001",
                    bullets=[
                        TailoredSentence(
                            section="work_experience",
                            sentence="负责后端服务开发。",
                            source_fact_ids=["fact_001"],
                        )
                    ],
                )
            ],
        )

        formal_resume = assemble_formal_resume(
            sample_jd(),
            resume,
            draft,
        )

        work = formal_resume.work_experiences[0]
        self.assertEqual(work.company, "示例科技")
        self.assertEqual(work.job_title, "后端工程师")
        self.assertEqual(work.start_date, "2023.06")
        self.assertEqual(work.end_date, "2025.06")
        self.assertEqual(work.bullets, ["负责后端服务开发。"])

    def test_formal_honor_award_preserves_parsed_metadata(self):
        resume = sample_resume()
        draft = TailoredResumeDraft(
            jd_id="jd_test",
            resume_id=resume.resume_id,
            honor_awards=[
                TailoredHonorAward(
                    honor_award_id="honor_001",
                    bullets=[
                        TailoredSentence(
                            section="honor_award",
                            sentence="毕业设计获评校级优秀。",
                            source_fact_ids=["fact_honor_001"],
                        )
                    ],
                )
            ],
        )

        formal_resume = assemble_formal_resume(
            sample_jd(),
            resume,
            draft,
        )

        honor = formal_resume.honor_awards[0]
        self.assertEqual(honor.name, "优秀毕业设计")
        self.assertEqual(honor.issuer, "示例大学")
        self.assertEqual(honor.date, "2023.06")
        self.assertEqual(honor.bullets, ["毕业设计获评校级优秀。"])

    def test_build_revises_after_audit_and_checks_again(self):
        client = AuditedTailoringClient()

        (
            _match_report,
            initial_draft,
            first_report,
            revised_draft,
            final_report,
        ) = build_tailored_resume(sample_jd(), sample_resume(), client=client)

        self.assertIn("百万用户", initial_draft.summary[0].sentence)
        self.assertEqual(
            first_report.checks[0].support_status,
            PARTIALLY_SUPPORTED,
        )
        self.assertEqual(
            revised_draft.summary[0].sentence,
            "使用 FastAPI 构建后端服务。",
        )
        self.assertEqual(final_report.checks[0].support_status, SUPPORTED)
        self.assertEqual(client.revision_calls, 1)
        self.assertEqual(client.fact_check_calls, 2)

    def test_initial_build_and_review_can_run_as_separate_stages(self):
        client = AuditedTailoringClient()
        jd = sample_jd()
        resume = sample_resume()

        match_report, initial_draft = build_initial_tailored_resume(
            jd,
            resume,
            client=client,
        )

        self.assertIn("百万用户", initial_draft.summary[0].sentence)
        self.assertEqual(client.fact_check_calls, 0)
        self.assertEqual(client.revision_calls, 0)

        first_report, revised_draft, final_report = review_tailored_resume(
            jd,
            resume,
            match_report,
            initial_draft,
            client=client,
        )

        self.assertEqual(
            first_report.checks[0].support_status,
            PARTIALLY_SUPPORTED,
        )
        self.assertEqual(
            revised_draft.summary[0].sentence,
            "使用 FastAPI 构建后端服务。",
        )
        self.assertEqual(final_report.checks[0].support_status, SUPPORTED)
        self.assertEqual(client.fact_check_calls, 2)
        self.assertEqual(client.revision_calls, 1)

    def test_staged_http_endpoints_return_initial_then_reviewed_resume(self):
        jd = sample_jd()
        resume = sample_resume()
        match_report = RequirementMatchReport(
            jd_id=jd.jd_id,
            resume_id=resume.resume_id,
        )
        draft = TailoredResumeDraft(
            jd_id=jd.jd_id,
            resume_id=resume.resume_id,
            summary=[
                TailoredSentence(
                    section="advantages",
                    sentence="使用 FastAPI 构建后端服务。",
                    source_fact_ids=["fact_001"],
                )
            ],
        )
        report = FactCheckReport(
            jd_id=jd.jd_id,
            resume_id=resume.resume_id,
        )
        client = TestClient(app)

        with patch(
            "app.api.tailoring.build_initial_tailored_resume",
            return_value=(match_report, draft),
        ):
            initial_response = client.post(
                "/tailoring/build/initial",
                json={
                    "jd": jd.model_dump(),
                    "resume": resume.model_dump(),
                },
            )

        self.assertEqual(initial_response.status_code, 200)
        self.assertEqual(
            initial_response.json()["formal_resume"]["advantages"],
            ["使用 FastAPI 构建后端服务。"],
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            stored = {}
            with (
                patch(
                    "app.api.tailoring.review_tailored_resume",
                    return_value=(report, draft, report),
                ),
                patch(
                    "app.services.tailored_resume_storage.list_tailored_resumes",
                    return_value=[],
                ),
                patch(
                    "app.services.tailored_resume_storage._write_tailored_resume",
                    side_effect=lambda item: stored.__setitem__(item.tailored_resume_id, item),
                ),
            ):
                review_response = client.post(
                    "/tailoring/build/review",
                    json={
                        "jd": jd.model_dump(),
                        "resume": resume.model_dump(),
                        "match_report": match_report.model_dump(),
                        "draft": draft.model_dump(),
                    },
                )

        self.assertEqual(review_response.status_code, 200)
        self.assertEqual(
            review_response.json()["formal_resume"]["advantages"],
            ["使用 FastAPI 构建后端服务。"],
        )
        self.assertTrue(
            review_response.json()["saved_resume"]["tailored_resume_id"]
        )

    def test_build_runs_one_extra_revision_when_first_fix_still_fails(self):
        client = TwoPassRevisionClient()

        result = build_tailored_resume(
            sample_jd(),
            sample_resume(),
            client=client,
        )

        self.assertEqual(client.revision_calls, 2)
        self.assertEqual(client.fact_check_calls, 3)
        self.assertEqual(result[-1].checks[0].support_status, SUPPORTED)

    def test_missing_audit_sentence_is_not_treated_as_supported(self):
        draft = TailoredResumeDraft(
            jd_id="jd_test",
            resume_id="resume_test",
            summary=[
                TailoredSentence(
                    section="summary",
                    sentence="使用 FastAPI 构建后端服务。",
                    source_fact_ids=["fact_001"],
                )
            ],
        )

        report = validate_fact_check_report(
            report=FactCheckReport(
                jd_id="jd_test",
                resume_id="resume_test",
            ),
            jd_id="jd_test",
            resume=sample_resume(),
            draft=draft,
        )

        self.assertEqual(len(report.checks), 1)
        self.assertEqual(
            report.checks[0].support_status,
            PARTIALLY_SUPPORTED,
        )
        self.assertIn("omitted", report.checks[0].issue)

    def test_missing_requirement_is_added_as_unknown(self):
        jd = sample_jd().model_copy(
            update={
                "requirements": [
                    *sample_jd().requirements,
                    JDRequirement(
                        requirement_id="req_002",
                        category="required_skill",
                        requirement_text="Kubernetes",
                        keywords=["Kubernetes"],
                    ),
                ]
            }
        )

        report = validate_match_report(
            RequirementMatchReport(
                jd_id=jd.jd_id,
                resume_id="resume_test",
                matches=[],
            ),
            jd,
            sample_resume(),
        )

        self.assertEqual(
            [match.requirement_id for match in report.matches],
            ["req_001", "req_002"],
        )
        self.assertTrue(
            all(match.match_status == "unknown" for match in report.matches)
        )

    def test_formal_resume_can_be_edited_confirmed_and_downloaded(self):
        jd = sample_jd()
        resume = sample_resume()
        draft = TailoredResumeDraft(
            jd_id=jd.jd_id,
            resume_id=resume.resume_id,
            headline=jd.job_title,
            summary=[
                TailoredSentence(
                    section="summary",
                    sentence="使用 FastAPI 构建后端服务。",
                    source_fact_ids=["fact_001"],
                )
            ],
        )
        formal_resume = assemble_formal_resume(jd, resume, draft)

        with tempfile.TemporaryDirectory() as temp_dir:
            data_dir = Path(temp_dir)
            stored = {}

            def write_item(item):
                stored[item.tailored_resume_id] = item

            def load_item(item_id):
                if item_id not in stored:
                    raise FileNotFoundError(item_id)
                return stored[item_id]

            with (
                patch(
                    "app.services.tailored_resume_storage.list_tailored_resumes",
                    return_value=[],
                ),
                patch(
                    "app.services.tailored_resume_storage._write_tailored_resume",
                    side_effect=write_item,
                ),
                patch(
                    "app.services.tailored_resume_storage.load_tailored_resume",
                    side_effect=load_item,
                ),
                patch(
                    "app.api.tailoring.load_tailored_resume",
                    side_effect=load_item,
                ),
                patch(
                    "app.services.docx_export.TAILORED_RESUME_EXPORT_DIR",
                    data_dir,
                ),
            ):
                saved = save_tailored_resume(
                    jd,
                    resume,
                    draft,
                    formal_resume=formal_resume,
                )
                edited = formal_resume.model_copy(
                    update={"advantages": ["用户确认后的个人优势。"]}
                )
                client = TestClient(app)

                edit_response = client.patch(
                    f"/tailoring/saved/{saved.tailored_resume_id}/content",
                    json=FormalResumeUpdateRequest(
                        formal_resume=edited
                    ).model_dump(),
                )
                self.assertEqual(edit_response.status_code, 200)
                self.assertEqual(
                    edit_response.json()["formal_resume"]["advantages"],
                    ["用户确认后的个人优势。"],
                )

                confirm_response = client.post(
                    f"/tailoring/saved/{saved.tailored_resume_id}/confirm",
                    json={"formal_resume": edited.model_dump()},
                )
                self.assertEqual(confirm_response.status_code, 200)
                self.assertEqual(confirm_response.json()["status"], "confirmed")

                download_response = client.get(
                    f"/tailoring/saved/{saved.tailored_resume_id}/docx"
                )
                self.assertEqual(download_response.status_code, 200)
                self.assertTrue(download_response.content.startswith(b"PK"))

                docx_path = data_dir / f"{saved.tailored_resume_id}.docx"
                document = Document(docx_path)
                document_text = "\n".join(
                    paragraph.text for paragraph in document.paragraphs
                )
                self.assertIn("用户确认后的个人优势。", document_text)
                self.assertIn("wang@example.com", document_text)
                self.assertIn("教育经历", document_text)
                self.assertIn("工作经历", document_text)
                self.assertIn("示例科技", document_text)
                self.assertIn("简历生成系统", document_text)
                self.assertIn("荣誉奖项", document_text)
                self.assertIn("优秀毕业设计", document_text)
                self.assertLess(
                    document_text.index("教育经历"),
                    document_text.index("个人优势"),
                )
                self.assertLess(
                    document_text.index("个人优势"),
                    document_text.index("工作经历"),
                )
                self.assertLess(
                    document_text.index("工作经历"),
                    document_text.index("项目经历"),
                )
                self.assertLess(
                    document_text.index("项目经历"),
                    document_text.index("荣誉奖项"),
                )

    def test_docx_input_is_supported_and_readable(self):
        self.assertIn(".pdf", SUPPORTED_EXTENSIONS)
        self.assertIn(".docx", SUPPORTED_EXTENSIONS)

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "resume.docx"
            document = Document()
            document.add_paragraph("Python 与 FastAPI 项目经历")
            document.save(path)

            self.assertEqual(
                read_document_text(path),
                "Python 与 FastAPI 项目经历",
            )


if __name__ == "__main__":
    unittest.main()
